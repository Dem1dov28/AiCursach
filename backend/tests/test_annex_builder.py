"""Annex plan + annex_builder (P2 stage 4)."""

from __future__ import annotations

import json
import struct
import zlib
from pathlib import Path
from unittest.mock import patch

from docx import Document

from backend.domain.document.annex_plan import (
    annex_plan_issues,
    normalize_annex_plan,
)
from backend.infrastructure.adapters.docx_images import insert_figures_as_appendices
from backend.infrastructure.langgraph.nodes.annex_builder import annex_builder_node
from tools.common.docx_stp import body, h1_center, setup_margins


def _png_bytes(width: int = 8, height: int = 8) -> bytes:
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * width for _ in range(height))
    compressed = zlib.compress(raw, 9)

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", compressed)
        + chunk(b"IEND", b"")
    )


def test_normalize_annex_plan_assigns_letters():
    plan = normalize_annex_plan(
        {
            "appendices": [
                {
                    "letter": "А",
                    "title": "Диаграммы",
                    "kind": "figures",
                    "figures": [
                        {"path": "a.png", "caption": "Рисунок А.1 – Use case"},
                    ],
                },
                {
                    "title": "Графики",
                    "kind": "figures",
                    "figures": [{"path": "b.png"}],
                },
            ]
        }
    )
    assert plan["appendices"][0]["letter"] == "А"
    assert plan["appendices"][1]["letter"] == "Б"
    assert "Рисунок Б.1" in plan["appendices"][1]["figures"][0]["caption"]
    assert not annex_plan_issues(plan)


def test_annex_plan_issues_empty():
    assert annex_plan_issues({"appendices": []})


def test_annex_builder_heuristic_fallback(tmp_path, monkeypatch):
    from backend.infrastructure.adapters import docx_images as di
    from backend.infrastructure.langgraph.nodes import annex_builder as ab

    proj = tmp_path / "CW"
    png = proj / "Материалы" / "Диаграммы" / "png" / "usecase.png"
    png.parent.mkdir(parents=True)
    png.write_bytes(_png_bytes())

    monkeypatch.setattr(di, "project_dir", lambda _n: proj)
    monkeypatch.setattr(ab, "REPO_ROOT", tmp_path)

    with (
        patch.object(ab, "uses_coursework_docx", return_value=True),
        patch.object(ab, "collect_project_pngs", return_value=[png]),
        patch.object(
            ab,
            "_llm_annex_plan",
            return_value=({"appendices": []}, {}, ["bad"]),
        ),
        patch.object(ab, "gate_after_llm", return_value={}),
        patch.object(
            ab,
            "build_heuristic_annex_plan",
            return_value={
                "appendices": [
                    {
                        "letter": "А",
                        "title": "Диаграммы",
                        "status": "обязательное",
                        "kind": "figures",
                        "figures": [
                            {
                                "path": str(png.relative_to(tmp_path)),
                                "caption": "Рисунок А.1 – Use case",
                            }
                        ],
                    }
                ]
            },
        ),
    ):
        result = annex_builder_node(
            {
                "work_type": "coursework",
                "needs_annexes": True,
                "project_name": "CW",
                "topic": "ИС",
            }
        )

    assert result["annex_plan_done"] is True
    plan = json.loads(result["annex_plan"])
    assert plan["appendices"][0]["letter"] == "А"


def test_insert_from_annex_plan(tmp_path):
    docx_path = tmp_path / "report.docx"
    doc = Document()
    setup_margins(doc.sections[0])
    h1_center(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    body(doc, "[1] Источник", indent=False)
    doc.save(str(docx_path))

    img = tmp_path / "Материалы" / "Диаграммы" / "png" / "usecase.png"
    img.parent.mkdir(parents=True)
    img.write_bytes(_png_bytes())

    plan = {
        "appendices": [
            {
                "letter": "А",
                "title": "Диаграммы UML",
                "status": "обязательное",
                "kind": "figures",
                "figures": [
                    {
                        "path": str(img),
                        "caption": "Рисунок А.1 – Диаграмма вариантов использования",
                    }
                ],
            }
        ]
    }
    n, msg = insert_figures_as_appendices(docx_path, annex_plan=plan)
    assert n == 1
    texts = " ".join(p.text for p in Document(str(docx_path)).paragraphs)
    assert "Диаграммы UML" in texts
    assert "Рисунок А.1 – Диаграмма вариантов использования" in texts
