"""Тесты копирования титульника из примера DOCX."""

from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Pt

from backend.infrastructure.adapters.student_info import StudentFields
from backend.infrastructure.adapters.title_from_example import (
    append_title_from_example,
    parse_title_example,
)


def _make_example_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Министерство образования Республики Беларусь")
    doc.add_paragraph("Факультет инженерно-экономический")
    doc.add_paragraph("ОТЧЕТ по лабораторной работе №4 на тему")
    p = doc.add_paragraph("Старая тема лабораторной")
    p.runs[0].bold = True
    doc.add_paragraph("Вариант 7")
    doc.add_paragraph("Выполнил\tстудент гр. 111111")
    doc.add_paragraph("\tИ. И. Иванов")
    doc.add_paragraph("Проверила\tМ. М. Петрова")
    doc.add_page_break()
    doc.add_paragraph("Цель")
    doc.save(path)


def test_copy_title_and_substitute_fields():
    with tempfile.TemporaryDirectory() as tmp:
        example = Path(tmp) / "example.docx"
        _make_example_docx(example)

        meta = parse_title_example(example)
        assert meta.fields.group == "111111"
        assert meta.fields.variant == "7"
        assert meta.topic_lines == ["Старая тема лабораторной"]

        out = Document()
        student = StudentFields(
            group="473601",
            student_name="И. С. Демидов",
            teacher_name="Е. А. Колопенько",
            variant="10",
            lab_number="4",
        )
        ok, msg = append_title_from_example(
            out,
            example,
            student,
            topic="Новая тема JavaScript",
        )
        assert ok, msg
        text = "\n".join(p.text for p in out.paragraphs)
        assert "473601" in text
        assert "И. С. Демидов" in text
        assert "Е. А. Колопенько" in text
        assert "Вариант 10" in text
        assert "Новая тема JavaScript" in text
        assert "111111" not in text
        assert "И. И. Иванов" not in text
        assert "Старая тема" not in text
        assert "Министерство образования" in text
