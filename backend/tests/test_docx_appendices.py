"""P2 stage 2: PNG → приложения DOCX."""

from __future__ import annotations

import struct
import zlib
from pathlib import Path

from docx import Document

from backend.infrastructure.adapters.docx_images import (
    allocate_appendix_letters,
    classify_pngs,
    insert_figures_as_appendices,
    used_appendix_letters,
)
from tools.common.docx_stp import appendix_start, body, h1_center, setup_margins


def _png_bytes(width: int = 8, height: int = 8) -> bytes:
    """Минимальный валидный PNG (>64 байт) без внешних зависимостей."""
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * width for _ in range(height))
    compressed = zlib.compress(raw, 9)

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", compressed)
        + chunk(b"IEND", b"")
    )


def test_classify_pngs_by_folder(tmp_path: Path):
    diag = tmp_path / "Материалы" / "Диаграммы" / "png" / "usecase.png"
    chart = tmp_path / "Материалы" / "Графики" / "fig_a.png"
    shot = tmp_path / "Материалы" / "Скриншоты" / "screen.png"
    for p in (diag, chart, shot):
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(b"x")

    groups = classify_pngs([diag, chart, shot])
    titles = [t for t, _ in groups]
    assert titles == ["Диаграммы", "Графики результатов", "Иллюстрации"]
    assert groups[0][1] == [diag]
    assert groups[1][1] == [chart]
    assert groups[2][1] == [shot]


def test_allocate_skips_used_letters():
    letters = allocate_appendix_letters({"А", "Б"}, 2)
    assert letters == ["В", "Г"]


def test_insert_figures_as_appendices(tmp_path: Path):
    docx_path = tmp_path / "report.docx"
    doc = Document()
    setup_margins(doc.sections[0])
    h1_center(doc, "ЗАКЛЮЧЕНИЕ")
    body(doc, "Текст заключения.")
    h1_center(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    body(doc, "[1] Источник", indent=False)
    doc.save(str(docx_path))

    png_dir = tmp_path / "Материалы" / "Диаграммы" / "png"
    png_dir.mkdir(parents=True)
    img = png_dir / "idef0_a0.png"
    data = _png_bytes()
    assert len(data) >= 64
    img.write_bytes(data)

    n, msg = insert_figures_as_appendices(docx_path, [img])
    assert n == 1
    assert "Приложения" in msg or "А(" in msg

    saved = Document(str(docx_path))
    texts = [p.text for p in saved.paragraphs if p.text.strip()]
    assert any("Приложение" in t and "А" in t for t in texts)
    assert any("Рисунок А.1" in t for t in texts)
    # Не перед заключением: заключение раньше приложения
    idx_zak = next(i for i, t in enumerate(texts) if "ЗАКЛЮЧЕНИЕ" in t)
    idx_app = next(i for i, t in enumerate(texts) if "Приложение" in t)
    assert idx_zak < idx_app


def test_insert_appendices_skips_existing_letter(tmp_path: Path):
    docx_path = tmp_path / "report.docx"
    doc = Document()
    setup_margins(doc.sections[0])
    appendix_start(doc, "А", "Листинг программы")
    body(doc, "code…")
    doc.save(str(docx_path))
    assert "А" in used_appendix_letters(Document(str(docx_path)))

    img = tmp_path / "Материалы" / "Графики" / "fig.png"
    img.parent.mkdir(parents=True)
    img.write_bytes(_png_bytes())

    n, msg = insert_figures_as_appendices(docx_path, [img])
    assert n == 1
    texts = " ".join(p.text for p in Document(str(docx_path)).paragraphs)
    assert "Приложение\u00a0Б" in texts or "Приложение Б" in texts.replace("\u00a0", " ")
    assert "Рисунок Б.1" in texts
