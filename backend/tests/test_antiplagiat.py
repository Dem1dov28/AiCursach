"""Antiplagiat node + appendix insert (P2 stage 5)."""

from __future__ import annotations

import struct
import zlib
from pathlib import Path
from unittest.mock import patch

from docx import Document

from backend.infrastructure.adapters.antiplagiat import (
    clamp_originality_pct,
    find_antiplagiat_screenshot,
    insert_antiplagiat_appendix,
)
from backend.infrastructure.langgraph.nodes.antiplagiat import antiplagiat_node
from tools.common.docx_stp import body, h1_center, setup_margins


def _png_bytes(width: int = 8, height: int = 8) -> bytes:
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * width for _ in range(height))
    compressed = zlib.compress(raw, 9)

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", compressed)
        + chunk(b"IEND", b"")
    )


def test_clamp_originality_pct():
    assert clamp_originality_pct(120) == 99
    assert clamp_originality_pct(10) == 50
    assert clamp_originality_pct("93") == 93


def test_find_antiplagiat_screenshot(tmp_path, monkeypatch):
    from backend.infrastructure.adapters import antiplagiat as ap

    monkeypatch.setattr(ap, "project_dir", lambda _n: tmp_path)
    shots = tmp_path / "Материалы" / "Скриншоты"
    shots.mkdir(parents=True)
    path = shots / "antiplagiat.png"
    path.write_bytes(_png_bytes())
    found = find_antiplagiat_screenshot("X")
    assert found == path


def test_antiplagiat_node_with_screenshot(tmp_path, monkeypatch):
    shot = tmp_path / "antiplagiat.png"
    shot.write_bytes(_png_bytes())
    with (
        patch(
            "backend.infrastructure.langgraph.nodes.antiplagiat.uses_coursework_docx",
            return_value=True,
        ),
        patch(
            "backend.infrastructure.langgraph.nodes.antiplagiat.find_antiplagiat_screenshot",
            return_value=shot,
        ),
        patch(
            "backend.infrastructure.langgraph.nodes.antiplagiat.gate_after_llm",
            return_value={},
        ),
    ):
        result = antiplagiat_node(
            {
                "work_type": "coursework",
                "needs_antiplagiat": True,
                "project_name": "CW",
                "originality_pct": 94,
            }
        )
    assert result["antiplagiat_done"] is True
    assert result["antiplagiat_screenshot"] == str(shot)
    assert result["annex_reserved_letters"] == ["А"]
    assert result["originality_pct"] == 94


def test_antiplagiat_node_skip_after_hitl():
    with (
        patch(
            "backend.infrastructure.langgraph.nodes.antiplagiat.uses_coursework_docx",
            return_value=True,
        ),
        patch(
            "backend.infrastructure.langgraph.nodes.antiplagiat.find_antiplagiat_screenshot",
            return_value=None,
        ),
        patch(
            "backend.infrastructure.langgraph.nodes.antiplagiat.gate_after_llm",
            return_value={},
        ),
    ):
        result = antiplagiat_node(
            {
                "work_type": "coursework",
                "needs_antiplagiat": True,
                "project_name": "CW",
                "clarification_log": [
                    {
                        "agent": "antiplagiat",
                        "answer": "Продолжить без скриншота",
                    }
                ],
            }
        )
    assert result["antiplagiat_done"] is True
    assert result["antiplagiat_screenshot"] == ""
    assert result["annex_reserved_letters"] == []


def test_insert_antiplagiat_appendix(tmp_path):
    docx_path = tmp_path / "report.docx"
    doc = Document()
    setup_margins(doc.sections[0])
    h1_center(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    body(doc, "[1] Источник", indent=False)
    doc.save(str(docx_path))

    shot = tmp_path / "antiplagiat.png"
    shot.write_bytes(_png_bytes())

    ok, msg = insert_antiplagiat_appendix(docx_path, shot, letter="А", originality_pct=91)
    assert ok
    texts = " ".join(p.text for p in Document(str(docx_path)).paragraphs)
    assert "Антиплагиат" in texts
    assert "Рисунок А.1" in texts
    assert "91%" in texts
