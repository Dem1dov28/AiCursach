"""Сборка docx лабораторной по СТП БГУИР."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

from backend.infrastructure.adapters.student_info import StudentFields
from backend.infrastructure.adapters.title_from_example import (
    append_bsuir_lab_title_page,
    append_title_from_example,
    infer_lab_number,
)
from backend.infrastructure.adapters.html_screenshots import collect_program_screenshots
from backend.core.paths import REPO_ROOT as ROOT, project_dir

_SCREENSHOT_CAPTIONS = {
    "01_form": "главная форма приложения",
    "01_index": "главная форма приложения",
    "01_main": "главная форма приложения",
    "02_result": "результаты обработки формы",
    "02_console": "вывод программы в консоли",
}


def _figure_caption(path: Path, index: int) -> str:
    stem = path.stem.lower()
    for key, label in _SCREENSHOT_CAPTIONS.items():
        if stem == key or stem.endswith(key):
            return f"Рисунок {index} – {label}"
    title = path.stem.replace("_", " ").replace("-", " ")
    return f"Рисунок {index} – {title}"


def _is_valid_image(path: Path) -> bool:
    if not path.is_file() or path.stat().st_size < 64:
        return False
    try:
        head = path.read_bytes()[:12]
    except OSError:
        return False
    if not (head.startswith(b"\x89PNG\r\n\x1a\n") or head[:3] == b"\xff\xd8\xff"):
        return False
    try:
        from PIL import Image

        with Image.open(path) as im:
            w, h = im.size
        if w < 200 or h < 200:
            return False
    except Exception:
        return False
    return True


def _import_stp():
    import sys

    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from tools.common.docx_stp import add_figure, body, h1, h1_center, setup_margins

    return add_figure, body, h1, h1_center, setup_margins


def _add_code_block(doc, text: str, *, keep_with_heading: bool = False) -> None:
    first_block = keep_with_heading
    for block in re.split(r"(?=// === )", text.strip()):
        block = block.strip()
        if not block:
            continue
        if block.startswith("// === "):
            title = block.split("\n", 1)[0].replace("// === ", "").replace(" ===", "").strip()
            p = doc.add_paragraph()
            run = p.add_run(title)
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
            if first_block:
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.widow_control = True
                first_block = False
            code_body = block.split("\n", 1)[1] if "\n" in block else ""
        else:
            code_body = block
        for line in code_body.splitlines() or [""]:
            p = doc.add_paragraph()
            run = p.add_run(line if line else " ")
            run.font.name = "Courier New"
            run.font.size = Pt(10)
            pf = p.paragraph_format
            pf.first_line_indent = Pt(0)
            pf.left_indent = Pt(12)
            pf.space_after = Pt(0)
            pf.widow_control = True
            if first_block:
                pf.keep_with_next = True
                first_block = False


def _add_section_body(doc, body_fn, text: str, *, lines_with_heading: int = 2) -> None:
    text = (text or "").strip()
    if not text:
        return
    attached = 0
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        for line in para.splitlines():
            line = line.strip()
            if line:
                body_fn(doc, line, keep_with_next=attached < lines_with_heading)
                attached += 1


def write_lab_docx_stp(
    project_name: str,
    lab_content: dict,
    *,
    student_info: dict | None = None,
    student_fields: StudentFields | None = None,
    example_docx: Path | None = None,
    screenshot_paths: list[Path] | None = None,
    console_output: str = "",
    code_language: str = "",
    code_run_success: bool | None = None,
) -> tuple[bool, str, Path]:
    add_figure, body, h1, h1_center, setup_margins = _import_stp()

    out_dir = project_dir(project_name) / "Отчет"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "Лабораторная_работа.docx"

    info = student_info or {}
    topic = lab_content.get("topic") or info.get("lab_title") or "Лабораторная работа"
    fields = student_fields or StudentFields()

    doc = Document()
    setup_margins(doc.sections[0])
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(14)

    if not fields.lab_number:
        fields.lab_number = infer_lab_number(
            topic,
            str(lab_content.get("topic") or ""),
            str(info.get("lab_title") or ""),
        )

    title_msg = ""
    if example_docx and example_docx.is_file():
        ok_title, title_msg = append_title_from_example(
            doc,
            example_docx,
            fields,
            topic=topic,
        )
        if not ok_title:
            append_bsuir_lab_title_page(doc, topic, fields)
            title_msg = f"{title_msg}; использован запасной шаблон"
    else:
        append_bsuir_lab_title_page(doc, topic, fields)
        title_msg = "Титульник: загрузите пример DOCX — будет скопирован один в один"

    sections: list[tuple[str, str, str]] = [
        ("Цель", lab_content.get("purpose", ""), "text"),
        ("Теоретические сведения", lab_content.get("theory", ""), "text"),
        ("Вариант задания", lab_content.get("variant_task", ""), "text"),
        ("Код программы", lab_content.get("program_code", ""), "code"),
        ("Работа программы", lab_content.get("program_work", ""), "work"),
        ("Выводы", lab_content.get("conclusions", ""), "text"),
    ]

    for title, content, kind in sections:
        h1(doc, title.upper(), keep_with_body=True)
        if kind == "code":
            code_text = str(content).strip()
            if code_text:
                _add_code_block(doc, code_text, keep_with_heading=True)
            else:
                body(doc, "Исходный код находится в папке src/ проекта.", keep_with_next=True)
        elif kind == "work":
            shots = [Path(p) for p in (screenshot_paths or []) if Path(p).exists()]
            if not shots:
                shots, _ = collect_program_screenshots(
                    project_name,
                    console_output=console_output,
                    code_language=code_language,
                    code_run_success=code_run_success,
                )

            work_text = str(content).strip()
            if shots and "скриншот" not in work_text.lower():
                if work_text:
                    work_text += "\n\n"
                work_text += "Ниже представлены скриншоты работы программы."
            elif not work_text and shots:
                work_text = "Ниже представлены скриншоты работы программы."

            _add_section_body(doc, body, work_text)

            if shots:
                for i, img in enumerate(shots, 1):
                    if _is_valid_image(img):
                        add_figure(doc, img, _figure_caption(img, i))
            elif code_run_success is False:
                body(
                    doc,
                    "Программа не была запущена из-за ошибок компиляции или выполнения. "
                    "Исправьте исходный код в папке src/ проекта и пересоберите отчёт.",
                    keep_with_next=True,
                )
            else:
                body(
                    doc,
                    "Скриншоты не сформированы. Для HTML нужен Chrome; "
                    "для консольных программ — успешный запуск с выводом в консоль.",
                    keep_with_next=True,
                )
        else:
            _add_section_body(doc, body, str(content))

    doc.save(out_path)
    msg = f"Сохранено: {out_path}"
    if title_msg:
        msg = f"{title_msg}. {msg}"
    return True, msg, out_path
