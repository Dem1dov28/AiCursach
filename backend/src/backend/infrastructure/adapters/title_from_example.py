"""Титульный лист: копирование из примера DOCX с подстановкой полей."""

from __future__ import annotations

import re
from copy import deepcopy
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from backend.infrastructure.adapters.student_info import StudentFields

FONT = "Times New Roman"
SIZE = Pt(14)
LINE = Pt(18)
TAB_COL = Cm(10.5)
KAFEDRA_TAB = Cm(3.2)

_TITLE_END = re.compile(
    r"^(?:"
    r"цель|"
    r"содержание|"
    r"реферат|"
    r"введение|"
    r"оглавление|"
    r"пояснительная записка"
    r")\s*$",
    re.I,
)
_GROUP_NUM = re.compile(
    r"(студент\s+гр\.?\s*|студент\s+группы\s*)([0-9]{3,6}[а-яА-Яa-zA-Z]?)",
    re.I,
)
_VARIANT_NUM = re.compile(r"(Вариант\s+)(\d+)", re.I)
_LAB_NUM_INLINE = re.compile(
    r"(лабораторн(?:ой|ая)\s+работ(?:ы|e)\s*№?\s*)(\d+)",
    re.I,
)
_FIO_RE = re.compile(
    r"[А-ЯЁ][а-яё]+\s+[А-ЯЁ]\.\s*[А-ЯЁ]\."
    r"|[А-ЯЁ]\.\s*[А-ЯЁ]\.\s*[А-ЯЁ][а-яё]+"
)
_FACULTY_RE = re.compile(r"Факультет\s+(.+)", re.I)
_KAFEDRA_RE = re.compile(r"Кафедра\s+(.+)", re.I)
_LAB_NUM_RE = re.compile(r"лабораторн(?:ой|ая)\s+работ(?:e|ы)\s*№?\s*(\d+)", re.I)
_TOPIC_START = re.compile(r"на\s+тему\s*$|на\s+тему", re.I)
_VARIANT_LINE = re.compile(r"^\s*вариант\s+\d+\s*$", re.I)
_SIGNATURE_START = re.compile(r"^\s*(?:выполнил|проверил|проверила)\b", re.I)


@dataclass
class TitleExampleMeta:
    fields: StudentFields = field(default_factory=StudentFields)
    topic_lines: list[str] = field(default_factory=list)
    paragraph_texts: list[str] = field(default_factory=list)


def _block_text(element) -> str:
    parts: list[str] = []
    for node in element.iter():
        if node.tag == qn("w:t") and node.text:
            parts.append(node.text)
    return "".join(parts).strip()


def _is_title_end(element) -> bool:
    tag = element.tag.split("}")[-1]
    if tag != "p":
        return False
    text = _block_text(element).strip()
    if not text:
        return False
    return bool(_TITLE_END.match(text))


def _is_paragraph(element) -> bool:
    return element.tag.split("}")[-1] == "p"


def _set_paragraph_text(paragraph_elem, text: str) -> None:
    nodes = [n for n in paragraph_elem.iter() if n.tag == qn("w:t")]
    if not nodes:
        return
    nodes[0].text = text
    for node in nodes[1:]:
        node.text = ""


def _split_topic_for_lines(topic: str, line_count: int) -> list[str]:
    topic = (topic or "").strip()
    if not topic:
        return [""] * max(line_count, 1)
    if line_count <= 1:
        return [topic]

    parts = [p.strip() for p in re.split(r"\n\s*\n", topic) if p.strip()]
    if len(parts) >= line_count:
        return parts[:line_count]

    sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", topic) if s.strip()]
    if len(sentences) >= line_count:
        chunk = max(1, len(sentences) // line_count)
        lines: list[str] = []
        idx = 0
        for i in range(line_count):
            if i == line_count - 1:
                lines.append(" ".join(sentences[idx:]))
            else:
                lines.append(" ".join(sentences[idx : idx + chunk]))
                idx += chunk
        return [ln for ln in lines if ln]

    while len(parts) < line_count:
        parts.append("")
    return parts[:line_count]


def infer_lab_number(*texts: str) -> str:
    for raw in texts:
        match = _LAB_NUM_RE.search(raw or "")
        if match:
            return match.group(1)
    return ""


def merge_title_fields(primary: StudentFields, fallback: StudentFields) -> StudentFields:
    return StudentFields(
        group=primary.group or fallback.group,
        student_name=primary.student_name or fallback.student_name,
        teacher_name=primary.teacher_name or fallback.teacher_name,
        variant=primary.variant or fallback.variant,
        faculty=primary.faculty or fallback.faculty,
        department=primary.department or fallback.department,
        lab_number=primary.lab_number or fallback.lab_number,
        checker_label=primary.checker_label or fallback.checker_label,
    )


def extract_title_fields_from_text(blob: str) -> StudentFields:
    fields = StudentFields()
    text = (blob or "").strip()
    if not text:
        return fields

    gm = _GROUP_NUM.search(text)
    if gm:
        fields.group = gm.group(2)

    vm = _VARIANT_NUM.search(text)
    if vm:
        fields.variant = vm.group(2)

    fm = _FACULTY_RE.search(text)
    if fm:
        fields.faculty = fm.group(1).strip()

    km = _KAFEDRA_RE.search(text)
    if km:
        fields.department = km.group(1).strip()

    fields.lab_number = infer_lab_number(text)

    if re.search(r"\bпроверила\b", text, re.I):
        fields.checker_label = "Проверила"
    elif re.search(r"\bпроверил\b", text, re.I):
        fields.checker_label = "Проверил"

    for match in _FIO_RE.finditer(text):
        name = match.group(0)
        ctx = text[max(0, match.start() - 80) : match.start()].lower()
        if any(k in ctx for k in ("проверил", "проверила", "руковод", "преподав")):
            if not fields.teacher_name:
                fields.teacher_name = name
        elif "студент" in ctx or not fields.student_name:
            if not fields.student_name:
                fields.student_name = name

    return fields


def _extract_topic_lines(paragraph_texts: list[str]) -> list[str]:
    topic_lines: list[str] = []
    after_report = False

    for raw in paragraph_texts:
        text = raw.strip()
        if not text:
            continue
        lower = text.lower()

        if _LAB_NUM_RE.search(text) or "отчет" in lower or "отчёт" in lower:
            after_report = True
            if _TOPIC_START.search(text) and not _TOPIC_START.search(text).group(0).endswith("тему"):
                tail = re.split(r"на\s+тему\s*", text, maxsplit=1, flags=re.I)
                if len(tail) == 2 and tail[1].strip():
                    topic_lines.append(tail[1].strip())
            continue

        if not after_report:
            continue

        if _VARIANT_LINE.match(text) or _SIGNATURE_START.match(text):
            break

        topic_lines.append(text)

    return topic_lines


def parse_title_example(example_path: Path) -> TitleExampleMeta:
    if not example_path.is_file():
        return TitleExampleMeta()

    doc = Document(example_path)
    paragraph_texts: list[str] = []
    for child in doc.element.body:
        tag = child.tag.split("}")[-1]
        if tag == "sectPr":
            continue
        if _is_title_end(child):
            break
        if tag == "p":
            paragraph_texts.append(_block_text(child))

    fields = extract_title_fields_from_text("\n".join(paragraph_texts))
    topic_lines = _extract_topic_lines(paragraph_texts)
    return TitleExampleMeta(
        fields=fields,
        topic_lines=topic_lines,
        paragraph_texts=paragraph_texts,
    )


def extract_example_fields(example_path: Path) -> StudentFields:
    return parse_title_example(example_path).fields


def _substitute_inline_text(text: str, example: TitleExampleMeta, merged: StudentFields) -> str:
    if not text:
        return text

    out = text

    if merged.group:
        out = _GROUP_NUM.sub(lambda m: f"{m.group(1)}{merged.group}", out)

    if merged.variant:
        out = _VARIANT_NUM.sub(lambda m: f"{m.group(1)}{merged.variant}", out)

    if merged.lab_number:
        out = _LAB_NUM_INLINE.sub(lambda m: f"{m.group(1)}{merged.lab_number}", out)

    ex = example.fields
    if ex.student_name and merged.student_name and ex.student_name != merged.student_name:
        out = out.replace(ex.student_name, merged.student_name)

    if ex.teacher_name and merged.teacher_name and ex.teacher_name != merged.teacher_name:
        out = out.replace(ex.teacher_name, merged.teacher_name)

    return out


def _substitute_text_nodes(element, example: TitleExampleMeta, merged: StudentFields) -> None:
    for node in element.iter():
        if node.tag != qn("w:t") or not node.text:
            continue
        node.text = _substitute_inline_text(node.text, example, merged)


def _substitute_topic_paragraphs(element, example: TitleExampleMeta, topic: str) -> None:
    if not topic or not example.topic_lines:
        return

    new_lines = _split_topic_for_lines(topic, len(example.topic_lines))
    old_lines = example.topic_lines

    for paragraph in element.iter():
        if not _is_paragraph(paragraph):
            continue
        current = _block_text(paragraph)
        if not current:
            continue
        for idx, old in enumerate(old_lines):
            if not old:
                continue
            if current == old or old in current or current in old:
                _set_paragraph_text(paragraph, new_lines[idx])
                break


def _apply_substitutions(
    element,
    example: TitleExampleMeta,
    merged: StudentFields,
    *,
    topic: str = "",
) -> None:
    _substitute_topic_paragraphs(element, example, topic)
    _substitute_text_nodes(element, example, merged)


def append_title_from_example(
    doc: Document,
    example_path: Path,
    student: StudentFields,
    *,
    topic: str = "",
) -> tuple[bool, str]:
    """Скопировать титульный лист из примера, подставив только изменяемые поля."""
    if not example_path.is_file():
        return False, f"Пример DOCX не найден: {example_path}"

    try:
        src = Document(example_path)
    except Exception as exc:
        return False, f"Не удалось открыть пример: {exc}"

    example = parse_title_example(example_path)
    merged = merge_title_fields(student, example.fields)

    if topic and not merged.lab_number:
        merged.lab_number = infer_lab_number(topic)

    copied = 0
    for child in list(src.element.body):
        tag = child.tag.split("}")[-1]
        if tag == "sectPr":
            continue
        if _is_title_end(child):
            break
        clone = deepcopy(child)
        _apply_substitutions(clone, example, merged, topic=topic)
        doc.element.body.append(clone)
        copied += 1

    if copied == 0:
        return False, "В примере не найден титульный лист"

    doc.add_page_break()
    return True, "Титульный лист скопирован из примера DOCX"


# --- запасной шаблон, если пример недоступен ---


def _run_fmt(run, *, bold: bool = False) -> None:
    run.font.name = FONT
    run.font.size = SIZE
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def _para_line_spacing(p) -> None:
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = LINE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def _spacer(doc: Document, lines: int = 1) -> None:
    for _ in range(lines):
        p = doc.add_paragraph()
        _para_line_spacing(p)


def _center(doc: Document, text: str, *, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run_fmt(p.add_run(text), bold=bold)
    _para_line_spacing(p)


def _left(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _run_fmt(p.add_run(text))
    _para_line_spacing(p)


def _left_tab(doc: Document, left: str, right: str = "", *, tab_pos=TAB_COL) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(tab_pos, WD_TAB_ALIGNMENT.LEFT)
    if left:
        _run_fmt(p.add_run(left))
    if right:
        if left:
            p.add_run("\t")
        _run_fmt(p.add_run(right))


def _tab_only(doc: Document, text: str, *, tab_pos=TAB_COL) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.tab_stops.add_tab_stop(tab_pos, WD_TAB_ALIGNMENT.LEFT)
    p.add_run("\t")
    _run_fmt(p.add_run(text))


def _checker_label(teacher_name: str) -> str:
    if not teacher_name:
        return "Проверил"
    token = teacher_name.strip().split()[-1].rstrip(".")
    if token.endswith(("а", "я")):
        return "Проверила"
    return "Проверил"


def append_bsuir_lab_title_page(doc: Document, topic: str, student: StudentFields) -> None:
    """Запасной титульный лист, если пример DOCX недоступен."""
    topic = (topic or "Лабораторная работа").strip()
    fields = student

    _center(doc, "Министерство образования Республики Беларусь")
    _center(doc, "Учреждение образования")
    _center(doc, "БЕЛОРУССКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ ИНФОРМАТИКИ И РАДИОЭЛЕКТРОНИКИ")

    if fields.faculty:
        fac = fields.faculty if fields.faculty.lower().startswith("факультет") else f"Факультет {fields.faculty}"
        _left(doc, fac)
    if fields.department:
        dept = re.sub(r"^кафедра\s+", "", fields.department, flags=re.I).strip()
        _left_tab(doc, "Кафедра", dept, tab_pos=KAFEDRA_TAB)

    _spacer(doc, 9)

    if fields.lab_number:
        _center(doc, f"ОТЧЕТ по лабораторной работе №{fields.lab_number} на тему")
    else:
        _center(doc, "ОТЧЕТ по лабораторной работе на тему")

    for part in re.split(r"\n\s*\n", topic):
        part = part.strip()
        if part:
            _center(doc, part, bold=True)

    if fields.variant:
        _spacer(doc, 3)
        _center(doc, f"Вариант {fields.variant}")

    _spacer(doc, 10)

    if fields.group or fields.student_name:
        _left_tab(doc, "Выполнил", f"студент гр. {fields.group}" if fields.group else "")
        if fields.student_name:
            _tab_only(doc, fields.student_name)

    if fields.teacher_name:
        label = fields.checker_label or _checker_label(fields.teacher_name)
        _left_tab(doc, label, fields.teacher_name)

    doc.add_page_break()


def append_simple_title_page(doc: Document, topic: str, student: StudentFields) -> None:
    append_bsuir_lab_title_page(doc, topic, student)
