"""Извлечение текста из загруженных файлов."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

from docx import Document


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts)


def _pdf_page_text(page) -> str:
    text = page.get_text("text") or page.get_text()
    return text.strip() if text else ""


def _pdf_page_ocr(page, *, language: str = "rus+eng") -> str:
    """OCR для сканов (нужен tesseract: brew install tesseract tesseract-lang)."""
    try:
        tp = page.get_textpage_ocr(language=language, dpi=200, full=True)
        text = page.get_text(textpage=tp) or ""
        return text.strip()
    except Exception:
        return ""


def _pdf_text(data: bytes, *, max_pages_ocr: int = 40) -> str:
    try:
        import fitz  # pymupdf
    except ImportError as exc:
        raise RuntimeError("Для PDF установите pymupdf: pip install pymupdf") from exc

    parts: list[str] = []
    with fitz.open(stream=data, filetype="pdf") as pdf:
        for page in pdf:
            text = _pdf_page_text(page)
            if text:
                parts.append(text)

        if not parts:
            for i, page in enumerate(pdf):
                if i >= max_pages_ocr:
                    parts.append(f"[… пропущено страниц после {max_pages_ocr}]")
                    break
                ocr = _pdf_page_ocr(page)
                if ocr:
                    parts.append(ocr)

    return "\n\n".join(parts)


def _rtf_text(data: bytes) -> str:
    raw = data.decode("utf-8", errors="replace")
    # грубое снятие RTF-разметки
    out = []
    skip = False
    for ch in raw:
        if ch == "{":
            skip = True
            continue
        if ch == "}":
            skip = False
            continue
        if not skip and ch.isprintable() or ch in "\n\t":
            out.append(ch)
    return "".join(out)


def extract_text_from_bytes(data: bytes, filename: str) -> str:
    if not data:
        raise ValueError(f"Файл «{filename}» пустой")

    name = (filename or "file.txt").lower()

    if name.endswith(".doc"):
        raise ValueError(
            f"Формат .doc не поддерживается («{filename}»). "
            "Сохраните файл как DOCX или PDF с текстовым слоем."
        )
    if name.endswith(".docx"):
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise ValueError(f"«{filename}» не похож на корректный DOCX")
        return _docx_text(data)
    if name.endswith(".txt") or name.endswith(".md"):
        return data.decode("utf-8", errors="replace")
    if name.endswith(".rtf"):
        return _rtf_text(data)
    if name.endswith(".pdf"):
        return _pdf_text(data)

    raise ValueError(
        f"Неподдерживаемый формат «{filename}». "
        "Используйте DOCX, PDF, TXT или MD."
    )


def extract_text_or_hint(data: bytes, filename: str) -> tuple[str, str | None]:
    """Вернуть (текст, подсказка_если_текст_пустой)."""
    text = extract_text_from_bytes(data, filename).strip()
    if text:
        return text, None

    name = (filename or "").lower()
    if name.endswith(".pdf"):
        hint = (
            f"Из PDF «{filename}» не извлечён текст (ни обычный, ни OCR). "
            "Установите: brew install tesseract tesseract-lang — или загрузите DOCX/TXT."
        )
    elif name.endswith(".docx"):
        hint = (
            f"DOCX «{filename}» не содержит читаемого текста "
            "(только картинки?). Добавьте текстовый слой или загрузите TXT."
        )
    else:
        hint = f"Файл «{filename}» не содержит текста."
    return "", hint


def extract_text_from_path(path: Path) -> str:
    return extract_text_from_bytes(path.read_bytes(), path.name)
