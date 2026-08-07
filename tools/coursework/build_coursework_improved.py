#!/usr/bin/env python3
"""Сборка пояснительной записки по СТП БГУИР 2024 из главы_md/*.md."""
from __future__ import annotations

import csv
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from calc_bz import STRATEGY_NAMES, best_strategy, compute_bz_strategies
from word_equations import add_formula_paragraph

BASE = Path(__file__).resolve().parent
MD_DIR = BASE / "главы_md"
IMG_DIR = BASE / "word_assets_final"
OUT_DOCX = BASE / "Курсовая_ИОУЗ_финальная_под_защиту.docx"
OUT_TXT = BASE / "Курсовая_ИОУЗ_финальная_под_защиту.txt"

FONT = "Times New Roman"
FONT_SIZE = Pt(14)
TABLE_FONT_SIZE = Pt(14)
APPENDIX_FONT_SIZE = Pt(14)
INDENT = Cm(1.25)
FIGURE_WIDTH = Cm(15)

THEME_LINE_1 = "«СОВЕРШЕНСТВОВАНИЕ СИСТЕМЫ УПРАВЛЕНИЯ МАТЕРИАЛЬНЫМИ ЗАПАСАМИ"
THEME_LINE_2 = "МАСТЕРСКОЙ ПО РЕМОНТУ КОМПЬЮТЕРНОЙ ТЕХНИКИ»"

CHAPTER_FILES = [
    "01_введение.md",
    "02_1.1.md",
    "03_1.2_литобзор.md",
    "04_1.3.md",
    "06_2.1.md",
    "07_2.2.md",
    "11_3.1.md",
    "12_3.2.md",
    "13_4_совершенствование.md",
    "14_3.4.md",
    "15_3.5.md",
    "16_5.3_капитал.md",
    "17_выводы.md",
    "18_заключение.md",
    "19_источники.md",
    "20_приложение_а.md",
    "21_приложение_б.md",
]

SECTION_INTROS = {
    "02_1.1.md": (
        "1 АНАЛИЗ ЛИТЕРАТУРНЫХ ИССЛЕДОВАНИЙ И ПРОГРАММНЫХ РЕШЕНИЙ",
        None,
    ),
    "06_2.1.md": ("2 МОДЕЛИРОВАНИЕ ПРЕДМЕТНОЙ ОБЛАСТИ", None),
    "11_3.1.md": (
        "3 КЛАССИФИКАЦИЯ НОМЕНКЛАТУР ПО ГРУППАМ",
        None,
    ),
    "13_4_совершенствование.md": (
        "4 СОВЕРШЕНСТВОВАНИЕ СИСТЕМЫ УПРАВЛЕНИЯ МАТЕРИАЛЬНЫМИ ЗАПАСАМИ",
        None,
    ),
    "14_3.4.md": ("5 ДОПОЛНИТЕЛЬНЫЕ ИССЛЕДОВАНИЯ", None),
}

PAGE_BREAK_BEFORE = {
    "01_введение.md",
    "02_1.1.md",
    "06_2.1.md",
    "11_3.1.md",
    "13_4_совершенствование.md",
    "14_3.4.md",
    "17_выводы.md",
    "18_заключение.md",
    "19_источники.md",
    "20_приложение_а.md",
}

CHAPTER_CONCLUSIONS = {
    "04_1.3.md": (
        "Таким образом, анализ предметной области, литературных подходов и программных "
        "средств обосновывает переход к формализованной системе управления запасами на "
        "основе ABC-XYZ-классификации и расчётных моделей оптимизации пополнения."
    ),
    "07_2.2.md": (
        "Таким образом, сформированная номенклатура и согласованная база исходных "
        "параметров обеспечивают достаточное качество данных для проведения "
        "классификационного и оптимизационного анализа."
    ),
    "12_3.2.md": (
        "Таким образом, для каждой группы матрицы ABC-XYZ определена стратегия "
        "управления запасами, соответствующая профилю спроса и экономической значимости "
        "номенклатур."
    ),
    "13_4_совершенствование.md": (
        "Таким образом, расчётные модели частичного и полного совмещения заказов, "
        "раздельной оптимизации и страхового запаса формируют набор альтернативных "
        "стратегий совершенствования системы управления запасами."
    ),
    "16_5.3_капитал.md": (
        "Таким образом, дополнительные исследования по группе BZ подтвердили "
        "экономическую целесообразность полного совмещения заказов, определили "
        "требуемую площадь склада 0,9892 м2 и показали различия стратегий по уровню "
        "капиталовложений в запасы."
    ),
}

UNNUMBERED_HEADINGS = {
    "РЕФЕРАТ",
    "СОДЕРЖАНИЕ",
    "ВВЕДЕНИЕ",
    "ВЫВОДЫ",
    "ЗАКЛЮЧЕНИЕ",
    "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ",
    "ПРИЛОЖЕНИЕ А",
    "ПРИЛОЖЕНИЕ Б",
}

TOC_ENTRIES = [
    ("Реферат", 0),
    ("Задание по курсовой работе", 0),
    ("Содержание", 0),
    ("Введение", 0),
    ("1 Анализ литературных исследований и программных решений", 0),
    ("1.1 Описание и анализ предметной области", 1),
    ("1.2 Обзор литературы по управлению запасами", 1),
    ("1.3 Анализ программных решений", 1),
    ("2 Моделирование предметной области", 0),
    ("2.1 Анализ и обоснованное формирование комплекса номенклатур", 1),
    ("2.2 Формирование исходных данных", 1),
    ("3 Классификация номенклатур по группам", 0),
    ("3.1 Описание проведения ABC-XYZ-анализа", 1),
    ("3.2 Определение стратегий управления запасами по группам", 1),
    ("4 Совершенствование системы управления материальными запасами", 0),
    ("4.1 Оптимизация по методу частичного совмещения заказов", 1),
    ("4.2 Оптимизация с учётом страхового запаса", 1),
    ("4.3 Оптимизация по методу полного совмещения заказов", 1),
    ("4.4 Оптимизация по методу раздельной оптимизации", 1),
    ("5 Дополнительные исследования", 0),
    ("5.1 Оптимизация группы BZ по различным стратегиям", 1),
    ("5.2 График занятости складских площадей", 1),
    ("5.3 График использования капиталовложений в запасы", 1),
    ("Выводы", 0),
    ("Заключение", 0),
    ("Список использованных источников", 0),
    (
        "Приложение А. Исходные данные и результаты ABC-XYZ-анализа",
        0,
    ),
    ("Приложение Б. График занятости складских площадей", 0),
]

SOURCE_HEADERS = [
    ("code", "Код"),
    ("name", "Наименование"),
    ("demand_v", "Спрос v, ед./сут."),
    ("coef_variation_cv", "CV"),
    ("unit_price_c", "Цена c, руб."),
    ("order_cost_K", "K, руб."),
    ("holding_cost_h", "h, руб./ед./сут."),
    ("storage_area_f", "f, м2/ед."),
]

ABC_HEADERS = [
    ("rank", "№"),
    ("code", "Код"),
    ("name", "Наименование"),
    ("annual_value", "Годовая стоимость"),
    ("share", "Доля"),
    ("cum_share", "Накопл. доля"),
    ("abc", "ABC"),
    ("xyz", "XYZ"),
    ("group", "Группа"),
]

FORBIDDEN_REPLACEMENTS = [
    (r"(?i)раздел\s+\d+\s+посвящен\s+", "Рассмотрен "),
    (r"(?i)раздел\s+\d+\s+открывает\s+", "Данный этап открывает "),
    (r"(?i)раздел\s+\d+\s+является\s+", "Данный этап является "),
    (r"(?i)\bраздел\s+\d+\b", "этап"),
    (r"(?i)по методическим (?:указаниям|требованиям)", "в рамках принятой методики"),
    (r"(?i)соответствует\s+методическим\s+требованиям", "соответствует принятой методике расчёта"),
    (r"(?i)в\s+соответствии\s+с\s+методическими\s+требованиями", "в рамках принятой методики"),
    (r"(?i)в\s+соответствии\s+с\s+индивидуальным\s+заданием", "в рамках поставленной задачи"),
    (r"(?i)в\s+соответствии\s+с\s+(?:заданием\s+)?вариант(?:ом|а)\s+\d+", "для выбранной группы номенклатур"),
    (r"(?i)индивидуальным\s+вариантом", "расчётной моделью"),
    (r"(?i)по\s+варианту\s+\d+", "для группы BZ"),
    (r"(?i)конкретного\s+варианта\s+задания", "расчётной модели мастерской"),
    (r"(?i)дополнительное\s+исследование\s+по\s+варианту\s+\d+", "дополнительное исследование группы BZ"),
    (r"(?i)результатах\s+предыдущего\s+раздела", "результатах предшествующего этапа анализа"),
]


def normalize_text(text: str) -> str:
    text = text.replace("\u2014", "-").replace("\u2013", "-").replace("–", "-").replace("—", "-")
    text = re.sub(r"\s+", " ", text).strip()
    for pattern, repl in FORBIDDEN_REPLACEMENTS:
        text = re.sub(pattern, repl, text, flags=re.IGNORECASE)
    text = re.sub(r"\s{2,}", " ", text)
    text = re.sub(r"^\s+,\s*", "", text)
    return text.strip()


def setup_paragraph(p, *, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=FONT_SIZE) -> None:
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = INDENT if indent else Cm(0)
    p.alignment = align
    for run in p.runs:
        run.font.name = FONT
        run.font.size = size


def setup_styles(doc: Document) -> None:
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = FONT_SIZE
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.line_spacing = 1.0
    pf.first_line_indent = INDENT
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    for s in doc.sections:
        s.left_margin = Cm(3)
        s.right_margin = Cm(1.5)
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)


def add_blank_line(doc: Document) -> None:
    p = doc.add_paragraph("")
    setup_paragraph(p, indent=False)


def add_runs(p, text: str, bold=False, italic=False) -> None:
    text = normalize_text(text)
    if "**" in text:
        parts = re.split(r"(\*\*.*?\*\*)", text)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                run = p.add_run(part[2:-2])
                run.bold = True
                run.italic = italic
            elif part:
                p.add_run(part)
    else:
        run = p.add_run(text)
        if bold:
            run.bold = True
        if italic:
            run.italic = italic
    for run in p.runs:
        run.font.name = FONT
        run.font.size = FONT_SIZE
        if bold:
            run.bold = True


def add_center_heading(doc: Document, text: str, txt: list[str]) -> None:
    add_blank_line(doc)
    p = doc.add_paragraph()
    add_runs(p, text, bold=True)
    setup_paragraph(p, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_blank_line(doc)
    txt.append(text)


def add_heading(doc: Document, text: str, txt: list[str]) -> None:
    text = normalize_text(text)
    if text.upper() in UNNUMBERED_HEADINGS or text in UNNUMBERED_HEADINGS:
        add_center_heading(doc, text.upper(), txt)
        return
    if re.match(r"^\d+\s+[А-ЯA-Z]", text):
        add_blank_line(doc)
        p = doc.add_paragraph()
        add_runs(p, text, bold=True)
        setup_paragraph(p, indent=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        add_blank_line(doc)
        txt.append(text)
        return
    if re.match(r"^\d+\.\d+", text):
        add_blank_line(doc)
        p = doc.add_paragraph()
        add_runs(p, text, bold=True)
        setup_paragraph(p, indent=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        add_blank_line(doc)
        txt.append(text)
        return
    add_center_heading(doc, text, txt)


def add_blank_formula(doc: Document) -> None:
    add_blank_line(doc)


def add_formula(doc: Document, text: str, txt: list[str]) -> None:
    add_blank_formula(doc)
    inner = text.strip().strip("*")
    if not add_formula_paragraph(doc, inner, txt):
        add_paragraph(doc, text, txt, formula=True)
    add_blank_formula(doc)


def add_paragraph(doc: Document, text: str, txt: list[str], *, center=False, formula=False) -> None:
    text = normalize_text(text)
    if not text:
        return
    p = doc.add_paragraph()
    add_runs(p, text, italic=formula)
    setup_paragraph(
        p,
        indent=not center,
        align=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    txt.append(text)


def add_list_items(doc: Document, items: list[str], txt: list[str]) -> None:
    for i, item in enumerate(items):
        ending = "." if i == len(items) - 1 else ";"
        line = f"- {normalize_text(item).rstrip('.;')}{ending}"
        p = doc.add_paragraph()
        p.add_run(line)
        setup_paragraph(p, indent=True)
        txt.append(line)


def read_tsv(name: str) -> list[dict[str, str]]:
    with (BASE / name).open(encoding="utf-8") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def style_table(table, size=TABLE_FONT_SIZE) -> None:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                setup_paragraph(p, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=size)
                for run in p.runs:
                    run.font.name = FONT
                    run.font.size = size


def add_table_caption(doc: Document, number: str, title: str, txt: list[str]) -> None:
    cap = f"Таблица {number} – {title}"
    p = doc.add_paragraph(cap)
    setup_paragraph(p, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    txt.append(cap)


def add_data_table(
    doc: Document,
    headers: list[tuple[str, str]],
    rows: list[dict[str, str]],
    txt: list[str],
    *,
    size=TABLE_FONT_SIZE,
) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, (_, title) in enumerate(headers):
        t.cell(0, i).text = title
    for row in rows:
        cells = t.add_row().cells
        for i, (key, _) in enumerate(headers):
            val = row.get(key, "")
            if key in {"annual_value", "share", "cum_share"}:
                try:
                    val = f"{float(val):.4f}" if key != "annual_value" else f"{float(val):.2f}"
                except ValueError:
                    pass
            cells[i].text = str(val)
    style_table(t, size=size)
    add_blank_line(doc)


def add_table_2_1(doc: Document, txt: list[str]) -> None:
    rows = read_tsv("приложение_А_исходные_данные_50.tsv")[:10]
    add_blank_line(doc)
    add_table_caption(doc, "2.1", "Исходные данные по номенклатурам (фрагмент)", txt)
    add_data_table(doc, SOURCE_HEADERS, rows, txt)


def add_matrix_table(doc: Document, txt: list[str]) -> None:
    rows = read_tsv("матрица_ABC_XYZ.tsv")
    add_blank_line(doc)
    add_table_caption(doc, "3.1", "Матрица ABC-XYZ", txt)
    headers = [("group", "Группа"), ("count", "Кол-во"), ("codes", "Коды номенклатур")]
    add_data_table(doc, headers, rows, txt)


def add_bz_table(doc: Document, txt: list[str]) -> None:
    add_blank_line(doc)
    add_table_caption(doc, "5.1", "Сравнение стратегий управления запасами для группы BZ", txt)
    strategies = compute_bz_strategies()
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    headers = ["Стратегия", "Суточные затраты, ден. ед.", "Годовые затраты, ден. ед."]
    for i, h in enumerate(headers):
        t.cell(0, i).text = h
    for name in STRATEGY_NAMES:
        daily, yearly = strategies[name]
        row = t.add_row().cells
        row[0].text = name
        row[1].text = f"{daily:.2f}"
        row[2].text = f"{yearly:.2f}"
    style_table(t)
    add_blank_line(doc)
    add_paragraph(
        doc,
        f"По результатам расчётов предпочтительной признана стратегия: {best_strategy(strategies)}.",
        txt,
    )


def add_appendix_tables(doc: Document, txt: list[str]) -> None:
    source = read_tsv("приложение_А_исходные_данные_50.tsv")
    abc = read_tsv("результаты_ABC_XYZ.tsv")
    matrix = read_tsv("матрица_ABC_XYZ.tsv")

    add_table_caption(doc, "А.1", "Исходные данные по 50 номенклатурам", txt)
    add_data_table(doc, SOURCE_HEADERS, source, txt, size=APPENDIX_FONT_SIZE)

    add_table_caption(doc, "А.2", "Результаты ABC-XYZ-анализа", txt)
    add_data_table(doc, ABC_HEADERS, abc, txt, size=APPENDIX_FONT_SIZE)

    add_table_caption(doc, "А.3", "Матрица ABC-XYZ", txt)
    matrix_headers = [("group", "Группа"), ("count", "Кол-во"), ("codes", "Коды номенклатур")]
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for i, (_, title) in enumerate(matrix_headers):
        t.cell(0, i).text = title
    for row in matrix:
        cells = t.add_row().cells
        cells[0].text = row["group"]
        cells[1].text = row["count"]
        cells[2].text = row["codes"]
    style_table(t, APPENDIX_FONT_SIZE)
    add_blank_line(doc)


def extract_backtick(text: str) -> str:
    if "`" in text:
        return text.split("`", 2)[1]
    return ""


def parse_image_block(lines: list[str], i: int) -> tuple[str, str, int] | None:
    line = lines[i].strip()
    if not line.startswith("**Вставить скриншот:**"):
        return None
    rel = extract_backtick(line)
    j = i + 1
    caption = ""
    while j < len(lines):
        nxt = lines[j].strip()
        if not nxt:
            j += 1
            continue
        if nxt.startswith("**Подпись:**"):
            caption = extract_backtick(nxt)
            j += 1
            break
        break
    return rel, caption, j


def setup_image_paragraph(p) -> None:
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.left_indent = Cm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.line_spacing = 1.0
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_picture(doc: Document, rel_path: str, caption: str, txt: list[str]) -> None:
    add_blank_line(doc)
    img = BASE / rel_path
    if not img.exists():
        alt = IMG_DIR / Path(rel_path).name
        img = alt if alt.exists() else img
    pic = doc.add_paragraph()
    setup_image_paragraph(pic)
    if img.exists():
        run = pic.add_run()
        run.add_picture(str(img), width=FIGURE_WIDTH)
    else:
        run = pic.add_run(f"[Изображение не найдено: {rel_path}]")
        run.font.name = FONT
        run.font.size = FONT_SIZE
    if caption:
        cap = normalize_text(caption)
        p = doc.add_paragraph(cap)
        setup_paragraph(p, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        txt.append(cap)
    add_blank_line(doc)


def add_toc(doc: Document, txt: list[str]) -> None:
    add_center_heading(doc, "СОДЕРЖАНИЕ", txt)
    for title, level in TOC_ENTRIES:
        p = doc.add_paragraph()
        indent_spaces = " " * (level * 2)
        dots = "." * max(3, 62 - len(title) - level * 2)
        p.add_run(f"{indent_spaces}{title}{dots}")
        setup_paragraph(p, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
        txt.append(title)
    doc.add_page_break()


def add_page_number_footer(section, start_at: int | None = None) -> None:
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if start_at is not None:
        sect_pr = section._sectPr
        pg_num_type = OxmlElement("w:pgNumType")
        pg_num_type.set(qn("w:start"), str(start_at))
        sect_pr.append(pg_num_type)
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_text)
    run._r.append(fld_end)
    run.font.name = FONT
    run.font.size = FONT_SIZE


def build_title_page(doc: Document, txt: list[str]) -> None:
    lines = [
        ("Министерство образования Республики Беларусь", False, WD_ALIGN_PARAGRAPH.CENTER),
        ("", False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Учреждение образования", False, WD_ALIGN_PARAGRAPH.CENTER),
        ("«Белорусский государственный университет", False, WD_ALIGN_PARAGRAPH.CENTER),
        ("информатики и радиоэлектроники»", False, WD_ALIGN_PARAGRAPH.CENTER),
        ("", False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Факультет инженерно-экономический", False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Кафедра экономической информатики", False, WD_ALIGN_PARAGRAPH.CENTER),
        ("", False, WD_ALIGN_PARAGRAPH.CENTER),
        ("«К ЗАЩИТЕ ДОПУСТИТЬ»", True, WD_ALIGN_PARAGRAPH.CENTER),
        ("", False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Руководитель курсового проекта", False, WD_ALIGN_PARAGRAPH.LEFT),
        ("________________ _________________", False, WD_ALIGN_PARAGRAPH.LEFT),
        ("", False, WD_ALIGN_PARAGRAPH.LEFT),
        ("", False, WD_ALIGN_PARAGRAPH.LEFT),
        ("ПОЯСНИТЕЛЬНАЯ ЗАПИСКА", True, WD_ALIGN_PARAGRAPH.CENTER),
        ("к курсовому проекту", False, WD_ALIGN_PARAGRAPH.CENTER),
        ("на тему:", False, WD_ALIGN_PARAGRAPH.CENTER),
        (THEME_LINE_1, True, WD_ALIGN_PARAGRAPH.CENTER),
        (THEME_LINE_2, True, WD_ALIGN_PARAGRAPH.CENTER),
        ("", False, WD_ALIGN_PARAGRAPH.CENTER),
        ("БГУИР КР 6-05-0611-01 ___ ПЗ", False, WD_ALIGN_PARAGRAPH.CENTER),
        ("", False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Выполнил студент группы ________", False, WD_ALIGN_PARAGRAPH.LEFT),
        ("_______________________________", False, WD_ALIGN_PARAGRAPH.LEFT),
        ("(подпись студента)", False, WD_ALIGN_PARAGRAPH.LEFT),
        ("", False, WD_ALIGN_PARAGRAPH.LEFT),
        ("Минск 2026", False, WD_ALIGN_PARAGRAPH.CENTER),
    ]
    for text, bold, align in lines:
        p = doc.add_paragraph()
        if text:
            add_runs(p, text, bold=bold)
        setup_paragraph(p, indent=False, align=align)
        if text:
            txt.append(text)
    doc.add_page_break()


def build_assignment_page(doc: Document, txt: list[str]) -> None:
    add_center_heading(doc, "ЗАДАНИЕ", txt)
    add_center_heading(doc, "на курсовую работу", txt)
    add_paragraph(
        doc,
        "по дисциплине «Инструментальное обеспечение управления запасами»",
        txt,
        center=True,
    )
    add_blank_line(doc)
    add_paragraph(doc, "1. Тема курсовой работы:", txt)
    add_paragraph(
        doc,
        f"{THEME_LINE_1} {THEME_LINE_2}",
        txt,
    )
    add_paragraph(doc, "2. Сроки сдачи студентом законченной работы: __________.", txt)
    add_paragraph(
        doc,
        "3. Исходные данные: моделирование системы управления запасами мастерской "
        "по ремонту компьютерной техники; номенклатура 50 позиций; ABC-XYZ-анализ; "
        "сравнение стратегий пополнения; анализ складских площадей и капиталовложений.",
        txt,
    )
    add_paragraph(
        doc,
        "4. Содержание пояснительной записки: титульный лист, реферат, задание, "
        "содержание, введение, разделы 1-5, выводы, заключение, список источников, приложения.",
        txt,
    )
    add_paragraph(
        doc,
        "5. Перечень графического материала: график занятости складских площадей; "
        "график использования капиталовложений в запасы.",
        txt,
    )
    add_paragraph(doc, "6. Консультант по курсовой работе: ___________________________.", txt)
    add_paragraph(doc, "7. Дата выдачи задания: __________.", txt)
    add_blank_line(doc)
    add_paragraph(doc, "Руководитель ____________________    Задание принял к исполнению __________", txt)


def parse_md_file(path: Path, doc: Document, txt: list[str], fname: str) -> None:
    if fname == "20_приложение_а.md":
        add_center_heading(doc, "ПРИЛОЖЕНИЕ А", txt)
        add_appendix_tables(doc, txt)
        return
    if fname == "21_приложение_б.md":
        add_center_heading(doc, "ПРИЛОЖЕНИЕ Б", txt)
        add_picture(
            doc,
            "word_assets_final/fig_5_2_warehouse_bz.png",
            "Рисунок Б.1 - График занятости складских площадей для группы BZ",
            txt,
        )
        return

    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    list_buf: list[str] = []

    def flush_list() -> None:
        nonlocal list_buf
        if list_buf:
            add_list_items(doc, list_buf, txt)
            list_buf = []

    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            flush_list()
            add_heading(doc, line[2:].strip(), txt)
            i += 1
            continue
        if line.startswith("**Вставить таблицу:**"):
            flush_list()
            key = extract_backtick(line)
            if key == "table_2_1":
                add_table_2_1(doc, txt)
            elif key == "table_matrix":
                add_matrix_table(doc, txt)
            i += 1
            continue
        if line.startswith("**Вставить скриншот:**"):
            flush_list()
            block = parse_image_block(lines, i)
            if block:
                rel, caption, i = block
                add_picture(doc, rel, caption, txt)
            continue
        if line.startswith("**Подпись:**"):
            i += 1
            continue
        if line.startswith("- "):
            list_buf.append(line[2:].strip())
            i += 1
            continue
        flush_list()
        if re.match(r"^\d+\)", line):
            add_list_items(doc, [line], txt)
            i += 1
            continue
        if fname == "14_3.4.md" and "Сравнительный расчёт" in line:
            add_paragraph(doc, line, txt)
            i += 1
            add_bz_table(doc, txt)
            continue
        if fname == "19_источники.md" and re.match(r"^\d+\.", line):
            p = doc.add_paragraph()
            add_runs(p, line)
            setup_paragraph(p, indent=True)
            txt.append(normalize_text(line))
            i += 1
            continue
        if line.startswith("**") and line.endswith("**"):
            add_formula(doc, line, txt)
            i += 1
            continue
        para = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "-", "**")):
            if re.match(r"^\d+\)", lines[i].strip()) or re.match(r"^\d+\.", lines[i].strip()):
                break
            para.append(lines[i].strip())
            i += 1
        add_paragraph(doc, " ".join(para), txt)
    flush_list()

    if fname in CHAPTER_CONCLUSIONS:
        add_paragraph(doc, CHAPTER_CONCLUSIONS[fname], txt)


def ensure_charts() -> None:
    script = BASE / "generate_charts.py"
    if script.exists():
        try:
            subprocess.run([sys.executable, str(script)], check=True, cwd=BASE)
        except subprocess.CalledProcessError as exc:
            print(f"Warning: chart generation failed: {exc}")


def main() -> None:
    ensure_charts()
    doc = Document()
    setup_styles(doc)
    txt_lines: list[str] = []

    # Титульник и реферат пользователь вставляет вручную перед этим файлом
    build_assignment_page(doc, txt_lines)
    doc.add_page_break()
    doc.add_page_break()
    add_toc(doc, txt_lines)

    doc.add_section()
    setup_styles(doc)
    add_page_number_footer(doc.sections[-1], start_at=1)

    for fname in CHAPTER_FILES:
        path = MD_DIR / fname
        if not path.exists():
            raise FileNotFoundError(path)

        if fname in PAGE_BREAK_BEFORE:
            doc.add_page_break()
            txt_lines.append("")

        if fname in SECTION_INTROS:
            title, intro = SECTION_INTROS[fname]
            add_heading(doc, title, txt_lines)
            if intro:
                add_paragraph(doc, intro, txt_lines)

        parse_md_file(path, doc, txt_lines, fname)

    doc.save(OUT_DOCX)
    OUT_TXT.write_text("\n".join(txt_lines) + "\n", encoding="utf-8")
    print(f"Created: {OUT_DOCX}")
    print(f"Created: {OUT_TXT}")


if __name__ == "__main__":
    main()
