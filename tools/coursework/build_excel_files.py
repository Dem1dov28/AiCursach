#!/usr/bin/env python3
"""Экспорт всех расчётных таблиц курсовой в отдельные Excel-файлы для сдачи."""
from __future__ import annotations

import shutil
from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent
MASTER = BASE / "таблицы_для_курсовой.xlsx"
OUT_DIR = BASE / "excel"

EXPORTS: list[tuple[str, str, str]] = [
    ("Source50", "01_Исходные_данные_50_номенклатур.xlsx", "раздел 2.3, приложение А"),
    ("ABC_XYZ", "02_ABC_XYZ_результаты_анализа.xlsx", "раздел 3.1"),
    ("Matrix", "03_Матрица_ABC_XYZ.xlsx", "раздел 3.2"),
    ("OptGroups", "04_Оптимизация_по_группам.xlsx", "раздел 3.3"),
    ("DetailEOQ", "05_Детальные_расчеты_EOQ.xlsx", "раздел 3.3"),
    ("BZ_Calc", "06_BZ_расчет_стратегий.xlsx", "раздел 3.4"),
    ("BZ_Compare", "07_BZ_сравнение_стратегий.xlsx", "раздел 3.4"),
    ("BZ_Graph12", "08_BZ_график_склад_12_периодов.xlsx", "разделы 3.5, 4, приложение Б"),
]


def export_sheet_xlsxwriter(src_wb, sheet_name: str, out_path: Path) -> None:
    """Копия листа через xlsxwriter (совместимо с Excel macOS)."""
    src = src_wb[sheet_name]
    wb = xlsxwriter.Workbook(
        str(out_path), {"strings_to_numbers": False, "use_future_functions": True}
    )
    ws = wb.add_worksheet(sheet_name[:31])
    for row in src.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            r, c = cell.row - 1, cell.column - 1
            if isinstance(cell.value, str) and cell.value.startswith("="):
                ws.write_formula(r, c, cell.value)
            elif isinstance(cell.value, bool):
                ws.write_boolean(r, c, cell.value)
            elif isinstance(cell.value, (int, float)):
                ws.write_number(r, c, cell.value)
            else:
                ws.write(r, c, str(cell.value))
    wb.close()


def write_manifest(created: list[tuple[str, str, str]]) -> None:
    lines = [
        "Электронные таблицы к курсовой работе по ИОУЗ",
        "Тема: совершенствование системы управления материальными запасами мастерской",
        "",
        "Сводный файл (все листы): excel/00_Все_таблицы_курсовой.xlsx",
        "",
        "Отдельные файлы:",
    ]
    for _, fname, section in created:
        lines.append(f"– {fname} ({section})")
    lines.extend(
        [
            "",
            "Таблицы содержат формулы Excel (как в лабораторных работах ИОУЗ).",
            "Файлы собраны через xlsxwriter для совместимости с Excel на macOS.",
            "Полный пересчёт всех формул — в сводном файле 00_Все_таблицы_курсовой.xlsx.",
            "",
            "Пересборка:",
            "  python3 build_excel_with_formulas.py",
            "  python3 build_excel_files.py",
        ]
    )
    (OUT_DIR / "Содержание_Excel_файлов.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    if not MASTER.exists():
        raise SystemExit(f"Не найден файл: {MASTER}")

    OUT_DIR.mkdir(exist_ok=True)
    wb = load_workbook(MASTER, data_only=False)

    master_copy = OUT_DIR / "00_Все_таблицы_курсовой.xlsx"
    shutil.copy2(MASTER, master_copy)

    created: list[tuple[str, str, str]] = []
    for sheet_name, filename, section in EXPORTS:
        if sheet_name not in wb.sheetnames:
            raise SystemExit(f"Лист «{sheet_name}» отсутствует в {MASTER.name}")
        out_path = OUT_DIR / filename
        export_sheet_xlsxwriter(wb, sheet_name, out_path)
        created.append((sheet_name, filename, section))
        print(f"Created: {out_path}")

    write_manifest(created)
    print(f"Created: {OUT_DIR / 'Содержание_Excel_файлов.txt'}")
    print(f"Copied:  {master_copy}")


if __name__ == "__main__":
    main()
