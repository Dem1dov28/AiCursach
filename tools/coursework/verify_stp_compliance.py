#!/usr/bin/env python3
"""Проверка финальной записки на соответствие замечаниям преподавателя."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING

BASE = Path(__file__).resolve().parent
DOC = BASE / "Курсовая_ИОУЗ_финальная_под_защиту.docx"


def main() -> int:
    doc = Document(DOC)
    text = "\n".join(p.text for p in doc.paragraphs)
    headings = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    fails: list[str] = []
    oks: list[str] = []

    forbidden = [
        (r"(?i)\bраздел\s+\d+\b", "упоминание «раздел N»"),
        (r"(?i)по\s+методическим\s+(?:указаниям|требованиям)", "«по методическим указаниям/требованиям»"),
        (r"(?i)соответствует\s+методическим\s+требованиям", "«соответствует методическим требованиям»"),
        (r"(?i)в\s+соответствии\s+с\s+вариант", "«в соответствии с вариантом»"),
        (r"(?i)индивидуальным\s+заданием", "«индивидуальным заданием»"),
        (r"(?i)конкретного\s+варианта\s+задания", "«конкретного варианта задания»"),
        (r"[\u2014\u2013]", "длинное/среднее тире"),
        (r"(?i)выводы\s+по\s+разделу", "подглава «выводы по разделу»"),
        (r"(?i)^1\.4\s+выводы", "подраздел 1.4"),
        (r"(?i)^2\.5\s", "подраздел 2.5"),
        (r"(?i)^3\.6\s", "подраздел 3.6"),
    ]
    for pat, name in forbidden:
        if re.search(pat, text):
            fails.append(name)

    structure = {
        "СОДЕРЖАНИЕ": "СОДЕРЖАНИЕ" in headings,
        "ВВЕДЕНИЕ": "ВВЕДЕНИЕ" in headings,
        "ВЫВОДЫ перед ЗАКЛЮЧЕНИЕМ": (
            "ВЫВОДЫ" in headings and "ЗАКЛЮЧЕНИЕ" in headings
            and headings.index("ВЫВОДЫ") < headings.index("ЗАКЛЮЧЕНИЕ")
        ),
        "1.2 литобзор": any("1.2 Обзор литературы" in h for h in headings),
        "Глава 4 (4.1-4.4)": any("4.1 Оптимизация" in h for h in headings),
        "Глава 5 (5.1-5.3)": any("5.1 Оптимизация группы BZ" in h for h in headings),
        "Приложение А полное имя в содержании": (
            "Приложение А. Исходные данные и результаты ABC-XYZ-анализа" in text
        ),
        "Приложение Б полное имя в содержании": (
            "Приложение Б. График занятости складских площадей" in text
        ),
    }
    for name, ok in structure.items():
        (oks if ok else fails).append(f"структура: {name}" if not ok else name)

    if len(doc.tables) >= 5:
        oks.append(f"таблицы Word: {len(doc.tables)}")
    else:
        fails.append(f"мало таблиц Word ({len(doc.tables)})")

    omml = len(doc._element.body.findall(".//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"))
    if omml >= 8:
        oks.append(f"формулы OMML: {omml}")
    else:
        fails.append(f"мало формул OMML ({omml})")

    figs = [p.text for p in doc.paragraphs if p.text.startswith("Рисунок")]
    if len(figs) >= 5:
        oks.append(f"рисунки: {len(figs)}")
    else:
        fails.append(f"мало рисунков ({len(figs)})")

    chapter_conclusions = sum(
        1 for p in doc.paragraphs if p.text.startswith("Таким образом,")
    )
    if chapter_conclusions >= 5:
        oks.append(f"выводы по главам (абзацы): {chapter_conclusions}")
    else:
        fails.append(f"выводов по главам мало ({chapter_conclusions}, нужно >=5)")

    non_single = sum(
        1
        for p in doc.paragraphs
        if p.paragraph_format.line_spacing_rule
        and p.paragraph_format.line_spacing_rule != WD_LINE_SPACING.SINGLE
    )
    if non_single == 0:
        oks.append("межстрочный интервал: одинарный")
    else:
        fails.append(f"не одинарный интервал в {non_single} абзацах")

    in_app_a = False
    app_a_text = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t == "ПРИЛОЖЕНИЕ А":
            in_app_a = True
            continue
        if in_app_a and t == "ПРИЛОЖЕНИЕ Б":
            break
        if in_app_a and t and not t.startswith("Таблица А."):
            if len(t) > 60 and not re.match(r"^\d", t):
                app_a_text.append(t)
    if not app_a_text:
        oks.append("приложение А: только таблицы")
    else:
        fails.append(f"пояснительный текст в приложении А ({len(app_a_text)} абз.)")

    print(f"Проверка: {DOC.name}\n")
    print("ИСПРАВЛЕНО / OK:")
    for x in oks:
        print(f"  [OK] {x}")
    if fails:
        print("\nОСТАЛОСЬ / FAIL:")
        for x in fails:
            print(f"  [!!] {x}")
    else:
        print("\nВсе автоматические проверки пройдены.")
    print(
        "\nПримечание: титул и реферат вставляются вручную перед заданием; "
        "при этом содержание окажется на 5-й странице (титул, реферат, задание, "
        "резервная страница, содержание)."
    )
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
