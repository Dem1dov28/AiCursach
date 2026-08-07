"""Оформление Word по СТП БГУИР: поля, абзацы, заголовки, рисунки, приложения."""

from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pathlib import Path

FONT = "Times New Roman"
SIZE = Pt(14)
INDENT = Cm(1.25)
FIGURE_WIDTH = Cm(15)


def setup_margins(section) -> None:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)


def add_run(p, text: str, *, bold: bool = False):
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = SIZE
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return run


def body_fmt(p, *, indent: bool = True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold: bool = False):
    pf = p.paragraph_format
    pf.first_line_indent = INDENT if indent else Cm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(18)
    pf.widow_control = True
    p.alignment = align
    for run in p.runs:
        run.font.name = FONT
        run.font.size = SIZE
        if bold:
            run.bold = True


def _keep_with_next(p) -> None:
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.widow_control = True


def body(doc: Document, text: str, *, indent: bool = True, keep_with_next: bool = False):
    p = doc.add_paragraph()
    add_run(p, text)
    body_fmt(p, indent=indent)
    if keep_with_next:
        _keep_with_next(p)
    return p


def blank_line(doc: Document, *, keep_with_next: bool = False):
    p = doc.add_paragraph()
    body_fmt(p, indent=False)
    if keep_with_next:
        _keep_with_next(p)
    return p


def h1_center(doc: Document, text: str):
    p = doc.add_paragraph()
    add_run(p, text, bold=True)
    body_fmt(p, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    blank_line(doc)
    return p


def h1(doc: Document, text: str, *, followed_by_subsection: bool = False, keep_with_body: bool = False):
    p = doc.add_paragraph()
    add_run(p, text, bold=True)
    body_fmt(p, indent=True, bold=True)
    p.paragraph_format.space_before = Pt(18)
    if keep_with_body:
        _keep_with_next(p)
        blank_line(doc, keep_with_next=True)
    elif not followed_by_subsection:
        blank_line(doc)
    else:
        _keep_with_next(p)
    return p


def h2(doc: Document, text: str, *, keep_with_body: bool = False):
    p = doc.add_paragraph()
    add_run(p, text, bold=True)
    body_fmt(p, indent=True, bold=True)
    if keep_with_body:
        _keep_with_next(p)
        blank_line(doc, keep_with_next=True)
    else:
        blank_line(doc)
    return p


def table_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    add_run(p, text)
    body_fmt(p, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    blank_line(doc)
    return p


def add_figure(doc: Document, image_path: Path | str, caption: str, *, width=FIGURE_WIDTH):
    blank_line(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    # Не применять body_fmt (EXACTLY Pt(18)) — иначе рисунок сжимается в одну строку
    p.add_run().add_picture(str(image_path), width=width)
    blank_line(doc)
    p2 = doc.add_paragraph()
    add_run(p2, caption)
    body_fmt(p2, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    blank_line(doc)


def appendix_start(doc: Document, letter: str, title: str, *, status: str = "обязательное"):
    if doc.paragraphs and doc.paragraphs[-1].text.strip():
        doc.add_page_break()
    center_line(doc, f"Приложение\u00a0{letter}")
    center_line(doc, f"({status})", bold=True)
    center_line(doc, title, bold=True)
    blank_line(doc)


def appendix_continuation(doc: Document, letter: str):
    doc.add_page_break()
    center_line(doc, f"Продолжение приложения {letter}", bold=True)
    blank_line(doc)


def center_line(doc: Document, text: str, *, bold: bool = False):
    p = doc.add_paragraph()
    add_run(p, text, bold=bold)
    body_fmt(p, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, bold=bold)
    return p


def add_page_number(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    for tag, val in (
        ("fldChar", "begin"),
        ("instrText", " PAGE "),
        ("fldChar", "separate"),
        ("t", "1"),
        ("fldChar", "end"),
    ):
        if tag == "instrText":
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = val
        elif tag == "t":
            node = OxmlElement("w:t")
            node.text = val
        else:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), val)
        run._r.append(node)
    run.font.name = FONT
    run.font.size = SIZE


class PageFlow:
  """Грубая оценка заполнения страницы для переносов (СТП ~2/3)."""

  def __init__(self, doc: Document, *, page_lines: int = 35):
      self.doc = doc
      self.page_lines = page_lines
      self.used = 0.0
      self.target = round(page_lines * 2 / 3)

  @staticmethod
  def text_lines(text: str) -> float:
      if not text:
          return 0.4
      return max(1.0, (len(text) + 85) / 86)

  def _break(self):
      self.doc.add_page_break()
      self.used = 0.0

  def force_new_page(self):
      if self.used > 0:
          self._break()
      else:
          self.used = 0.0

  def before(self, need: float):
      if self.used <= 0:
          return
      if need > self.page_lines - self.used:
          self._break()
      elif self.used >= self.target and need >= 1.0:
          self._break()

  def add_text(self, text: str):
      self.before(self.text_lines(text))
      self.used += self.text_lines(text)
