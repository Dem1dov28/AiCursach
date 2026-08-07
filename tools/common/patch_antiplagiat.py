#!/usr/bin/env python3
"""Приложение А — скриншот «Антиплагиат»; листинг → Б, SQL → В (как ПСП)."""

from __future__ import annotations

import argparse
import re
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

FONT = "Times New Roman"
APP_A_TITLE = "Отчет о проверке на заимствования в системе «Антиплагиат»"
FIG_CAPTION = "Рисунок А.1 – Отчет о проверке на заимствования в системе «Антиплагиат»"
APP_B_TITLE = "Листинг программы (фрагменты)"
APP_C_TITLE = "SQL-скрипт создания базы данных"


def _norm(text: str) -> str:
    return text.strip().replace("\xa0", " ")


def _is_hdr(text: str, letter: str) -> bool:
    return _norm(text) == f"Приложение {letter}"


def _find_last(paras, letter: str) -> int | None:
    idx = None
    for i, p in enumerate(paras):
        if _is_hdr(p.text, letter):
            idx = i
    return idx


def _delete(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def _after(paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def _page_break_after(paragraph) -> Paragraph:
    p = _after(paragraph)
    p_pr = p._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:pageBreakBefore"))
    return p


def _center(paragraph, text: str = "", *, bold: bool = False) -> None:
    paragraph.clear()
    if text:
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.name = FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.first_line_indent = Cm(0)


def _hdr_after(anchor: Paragraph, letter: str, title: str) -> Paragraph:
    p = _page_break_after(anchor)
    _center(p, f"Приложение\u00a0{letter}")
    p = _after(p)
    _center(p, "(обязательное)", bold=True)
    p = _after(p)
    _center(p, title, bold=True)
    p = _after(p)
    _center(p, "")
    return p


def patch(doc_path: Path, screenshot: Path) -> int:
    if not screenshot.exists():
        print(f"Нет скриншота: {screenshot}", file=sys.stderr)
        return 1
    if not doc_path.exists():
        print(f"Нет файла: {doc_path}", file=sys.stderr)
        return 1

    doc = Document(str(doc_path))
    paras = doc.paragraphs
    idx_a = _find_last(paras, "А")
    idx_b = _find_last(paras, "Б")
    if idx_a is None or idx_b is None or idx_b <= idx_a:
        print("Не найдены приложения А и Б", file=sys.stderr)
        return 1

    paras[idx_a + 2].text = APP_A_TITLE
    listing = [deepcopy(p._element) for p in paras[idx_a + 4 : idx_b]]
    for elem in listing:
        for node in elem.iter():
            if node.tag == qn("w:t") and node.text and "Продолжение приложения А" in node.text:
                node.text = node.text.replace("Продолжение приложения А", "Продолжение приложения Б")

    sql_start = idx_b + 4
    while sql_start < len(paras) and not paras[sql_start].text.strip():
        sql_start += 1
    sql_elems = [deepcopy(p._element) for p in paras[sql_start:]]

    for p in reversed(paras[idx_a + 4 :]):
        _delete(p)

    paras = doc.paragraphs
    anchor = paras[idx_a + 3]
    img_p = _after(anchor)
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.add_run().add_picture(str(screenshot), width=Cm(15))
    cap = _after(_after(img_p))
    _center(cap, FIG_CAPTION)

    b_hdr = _hdr_after(cap, "Б", APP_B_TITLE)
    after_b = anchor
    anchor_elem = b_hdr._p
    for elem in listing:
        anchor_elem.addnext(elem)
        after_b = Paragraph(elem, b_hdr._parent)
        anchor_elem = elem

    c_hdr = _hdr_after(after_b, "В", APP_C_TITLE)
    anchor_elem = c_hdr._p
    for elem in sql_elems:
        anchor_elem.addnext(elem)
        anchor_elem = elem

    for p in doc.paragraphs:
        t = p.text
        if "2 приложени" in t.lower():
            p.text = re.sub(r"2\s+приложени", "3 приложени", t, flags=re.I)
        if "двух приложений" in t:
            p.text = t.replace("двух приложений", "трёх приложений")

    doc.save(str(doc_path))
    print(f"OK: {doc_path}")
    return 0


def main():
    ap = argparse.ArgumentParser(description="Вставить приложение А (антиплагиат)")
    ap.add_argument("docx", type=Path, help="Пояснительная записка .docx")
    ap.add_argument("--screenshot", "-s", type=Path, required=True, help="PNG скриншота антиплагиата")
    args = ap.parse_args()
    raise SystemExit(patch(args.docx, args.screenshot))


if __name__ == "__main__":
    main()
