#!/usr/bin/env python3
"""Пересчёт номеров страниц в содержании (редактируйте HEADING_PATTERNS под проект)."""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document

DOCX_DEFAULT: Path | None = None

ORIGINALITY_PCT = 95
INTRO_PAGE_DEFAULT = 5
TOTAL_PAGES_DEFAULT = 67
LINES_PER_PAGE = 27

HEADING_PATTERNS: list[tuple[str, str, int]] = [
    (r"^ВВЕДЕНИЕ$", "Введение", 0),
    (
        r"^1 АНАЛИЗ",
        "1 Анализ литературных источников и программных решений",
        0,
    ),
    (r"^1\.1 ", "1.1 Описание и анализ предметной области", 1),
    (r"^1\.2 ", "1.2 Обзор аналогов программных средств", 1),
    (
        r"^2 МОДЕЛИРОВАНИЕ",
        "2 Моделирование предметной области и разработка требований",
        0,
    ),
    (r"^2\.1 ", "2.1 Бизнес-процессы и участники", 1),
    (r"^2\.2 ", "2.2 Функциональные и нефункциональные требования", 1),
    (r"^2\.3 ", "2.3 Образ предлагаемого решения", 1),
    (
        r"^3 ПРОЕКТИРОВАНИЕ",
        "3 Проектирование и разработка программного средства",
        0,
    ),
    (r"^3\.1 ", "3.1 Архитектура и технологии", 1),
    (r"^3\.2 ", "3.2 Проектирование интерфейса", 1),
    (r"^3\.3 ", "3.3 Модель данных", 1),
    (
        r"^3\.4 ",
        "3.4 Диаграммы классов серверной и клиентской части",
        1,
    ),
    (r"^3\.5 ", "3.5 Диаграммы поведения системы", 1),
    (r"^3\.6 ", "3.6 Алгоритмы бизнес-логики", 1),
    (r"^4 ТЕСТИРОВАНИЕ", "4 Тестирование и проверка работоспособности", 0),
    (r"^5 РУКОВОДСТВО", "5 Руководство по установке и использованию", 0),
    (r"^5\.1 ", "5.1 Установка", 1),
    (r"^5\.2 ", "5.2 Использование", 1),
    (r"^ЗАКЛЮЧЕНИЕ$", "Заключение", 0),
    (r"^СПИСОК ИСПОЛЬЗОВАННЫХ", "Список использованных источников", 0),
    (
        r"^Приложение[\s\xa0]А$",
        "Приложение А (обязательное) Отчет о проверке на заимствования в системе «Антиплагиат»",
        0,
    ),
    (
        r"^Приложение[\s\xa0]Б$",
        "Приложение Б (обязательное) Листинг серверной и клиентской части (фрагменты)",
        0,
    ),
    (
        r"^Приложение[\s\xa0]В$",
        "Приложение В (обязательное) SQL-скрипт создания базы данных",
        0,
    ),
]


def _est_lines(paragraph) -> float:
    text = paragraph.text.strip()
    if not text:
        return 0.35
    if paragraph._element.xpath(".//wp:inline"):
        return 15.0
    if paragraph._element.xpath(".//w:tbl"):
        return 11.0
    if paragraph.runs and paragraph.runs[0].font.name == "Courier New":
        size = paragraph.runs[0].font.size.pt if paragraph.runs[0].font.size else 10
        return 0.38 if size <= 9 else 0.48
    if text.startswith("Продолжение приложения"):
        return 2.5
    if text.isupper() and len(text) < 90:
        return 2.0
    if text.startswith("Таблица ") or text.startswith("Рисунок "):
        return 1.2
    return max(1.0, len(text) / 80)


def _find_body_start(doc: Document) -> int:
    seen_toc = False
    for i, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == "СОДЕРЖАНИЕ":
            seen_toc = True
        if seen_toc and paragraph.text.strip().upper() == "ВВЕДЕНИЕ":
            return i
    raise ValueError("Не найдено начало основного текста (ВВЕДЕНИЕ)")


def compute_toc_pages(
    doc: Document,
    *,
    intro_page: int,
    total_pages: int,
) -> list[tuple[str, int, int]]:
    body_start = _find_body_start(doc)
    page = intro_page
    fill = 0.0
    raw: dict[str, int] = {}

    for i in range(body_start, len(doc.paragraphs)):
        paragraph = doc.paragraphs[i]
        text = paragraph.text.strip().replace("\xa0", " ")
        for pattern, title, _level in HEADING_PATTERNS:
            if re.match(pattern, text, re.I) and title not in raw:
                raw[title] = page
        fill += _est_lines(paragraph)
        while fill >= LINES_PER_PAGE:
            fill -= LINES_PER_PAGE
            page += 1

    raw_end = page + (1 if fill > 4 else 0)
    span = max(1, raw_end - intro_page)
    scale = (total_pages - intro_page) / span

    entries: list[tuple[str, int, int]] = []
    prev = intro_page
    for _pattern, title, level in HEADING_PATTERNS:
        raw_page = raw.get(title, prev)
        if title == "Введение":
            scaled = intro_page
        else:
            scaled = intro_page + round((raw_page - intro_page) * scale)
        scaled = max(prev, scaled)
        entries.append((title, level, scaled))
        prev = scaled

    # Приложение Б — минимум на страницу после А (антиплагиат ~1 с.)
    for i, (title, level, page) in enumerate(entries):
        if title.startswith("Приложение Б") and i > 0:
            prev_page = entries[i - 1][2]
            if page <= prev_page:
                entries[i] = (title, level, prev_page + 1)

    return entries


def _count_stats(doc: Document) -> tuple[int, int, int]:
    figures = sum(
        1 for p in doc.paragraphs if re.match(r"^Рисунок \d+\.\d+", p.text.strip())
    )
    tables = sum(
        1 for p in doc.paragraphs if re.match(r"^Таблица \d+\.\d+", p.text.strip())
    )
    sources = sum(1 for p in doc.paragraphs if re.match(r"^\[\d+\]", p.text.strip()))
    return figures, tables, sources


def _set_toc_paragraph(paragraph, title: str, level: int, page: int) -> None:
    indent = "\u00a0\u00a0" * level
    paragraph.text = f"{indent}{title}\t{page}"


def _find_toc_range(doc: Document) -> tuple[int, int]:
    start = end = None
    for i, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == "СОДЕРЖАНИЕ":
            start = i
        if start is not None and i > start and paragraph.text.strip().upper() == "ВВЕДЕНИЕ":
            end = i
            break
    if start is None or end is None:
        raise ValueError("Не найден блок содержания")
    return start, end


def update(
    doc_path: Path,
    *,
    intro_page: int = INTRO_PAGE_DEFAULT,
    total_pages: int = TOTAL_PAGES_DEFAULT,
) -> int:
    if not doc_path.exists():
        print(f"Файл не найден: {doc_path}", file=sys.stderr)
        return 1

    doc = Document(str(doc_path))
    figures, tables, sources = _count_stats(doc)
    toc_entries = compute_toc_pages(
        doc, intro_page=intro_page, total_pages=total_pages
    )

    referat_stats = (
        f"Пояснительная записка {total_pages} с., {figures} рис., {tables} табл., "
        f"{sources} источников, 3 приложения."
    )
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Пояснительная записка") and "с." in text and "рис." in text:
            if paragraph.runs:
                paragraph.runs[0].text = referat_stats
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.text = referat_stats
            print("Реферат:", referat_stats)
        if "двух приложений" in text:
            paragraph.text = text.replace("двух приложений", "трёх приложений")

    toc_start, toc_end = _find_toc_range(doc)
    toc_lines = [
        p for p in doc.paragraphs[toc_start + 1 : toc_end] if "\t" in p.text
    ]

    if len(toc_lines) != len(toc_entries):
        print(
            f"Предупреждение: строк содержания {len(toc_lines)}, "
            f"ожидалось {len(toc_entries)}",
            file=sys.stderr,
        )

    for para, (title, level, page) in zip(toc_lines, toc_entries):
        _set_toc_paragraph(para, title, level, page)

    print(f"Содержание (введение — стр. {intro_page}, всего {total_pages} с.):")
    for title, level, page in toc_entries:
        prefix = "  " * level
        print(f"  {prefix}{title[:55]:55} {page}")

    doc.save(str(doc_path))
    print(f"Сохранено: {doc_path}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Обновить реферат и содержание")
    parser.add_argument("docx", type=Path, help="Пояснительная_записка.docx")
    parser.add_argument("--intro", type=int, default=INTRO_PAGE_DEFAULT)
    parser.add_argument("--total", type=int, default=TOTAL_PAGES_DEFAULT)
    args = parser.parse_args()
    return update(
        Path(args.docx),
        intro_page=args.intro,
        total_pages=args.total,
    )


if __name__ == "__main__":
    raise SystemExit(main())
