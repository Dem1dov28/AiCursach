"""Титульный лист и оглавление отчёта (ПСП_Ивановская + СТП 01-2024)."""

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Times New Roman"
SIZE = Pt(14)


def set_run_font(run, name=FONT, size=14, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_paragraph(
    doc,
    text="",
    *,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    bold=False,
    space_before=0,
    space_after=0,
):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if text:
        run = p.add_run(text)
        set_run_font(run, bold=bold)
    return p


def add_empty_lines(doc, count=1, *, align=WD_ALIGN_PARAGRAPH.CENTER):
    for _ in range(count):
        add_paragraph(doc, align=align)


def _set_table_borders_none(table):
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tbl_pr)
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "none")
        elem.set(qn("w:sz"), "0")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), "auto")
        borders.append(elem)
    tbl_pr.append(borders)


def _set_cell_text(cell, lines: list[tuple[str, WD_ALIGN_PARAGRAPH | None, bool]]):
    cell.text = ""
    for i, (text, align, bold) in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        if text:
            run = p.add_run(text)
            set_run_font(run, bold=bold)


def _signature_table(doc, right_lines: list[tuple[str, WD_ALIGN_PARAGRAPH | None, bool]]):
    """Таблица 1×2 без рамок: слева пусто, справа блок подписей (как в ПСП_Ивановская)."""
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    _set_table_borders_none(table)
    _set_cell_text(table.rows[0].cells[0], [("", WD_ALIGN_PARAGRAPH.LEFT, False)])
    _set_cell_text(table.rows[0].cells[1], right_lines)
    return table


def build_title_page(
    doc,
    topic: str,
    *,
    doc_code: str = "БГУИР КП 6-05-0611-01 010 ПЗ",
    discipline: str | None = "Программирование сетевых приложений",
    variant: int | None = 10,
    student_group: str = "473601",
    student_name: str = "И. С. Демидов",
    teacher_name: str = "Е.И. Пономарева",
    teacher_position: str = "преподаватель",
    city_year: str = "Минск 2026",
    lab_number: str = "",
    kind: str = "psp",
):
    """
    Титульный лист по образцу ПСП_Ивановская:
    шапка по центру, таблицы подписей без рамок, тема в «ёлочках», город/год внизу.
    kind='psp' — как в примере (курсовой проект); kind='lab' — отчёт по лабораторной.
    """
    add_paragraph(doc, "Министерство образования Республики Беларусь")
    add_paragraph(
        doc,
        "Учреждение образования «Белорусский государственный университет "
        "информатики и радиоэлектроники»",
    )
    add_empty_lines(doc, 1)
    add_paragraph(doc, "Факультет инженерно-экономический")
    add_paragraph(doc, "Кафедра экономической информатики")
    if discipline:
        add_paragraph(doc, f"Дисциплина «{discipline}»")

    add_empty_lines(doc, 4, align=WD_ALIGN_PARAGRAPH.LEFT)

    if kind == "lab":
        _signature_table(
            doc,
            [
                ("Проверила:", WD_ALIGN_PARAGRAPH.LEFT, False),
                (teacher_position, WD_ALIGN_PARAGRAPH.LEFT, False),
                (f"______________ {teacher_name}", WD_ALIGN_PARAGRAPH.LEFT, False),
                ("___.___.2026", WD_ALIGN_PARAGRAPH.LEFT, False),
            ],
        )
    else:
        _signature_table(
            doc,
            [
                ("«К защите допустить»", WD_ALIGN_PARAGRAPH.LEFT, False),
                ("Руководитель курсового проекта", WD_ALIGN_PARAGRAPH.LEFT, False),
                (
                    f"{teacher_position} ______________ {teacher_name}",
                    WD_ALIGN_PARAGRAPH.LEFT,
                    False,
                ),
                ("___.___.2026", WD_ALIGN_PARAGRAPH.LEFT, False),
            ],
        )

    add_empty_lines(doc, 4, align=WD_ALIGN_PARAGRAPH.LEFT)

    if kind == "lab":
        if lab_number:
            add_paragraph(doc, f"ОТЧЁТ по лабораторной работе №{lab_number}", bold=True)
        else:
            add_paragraph(doc, "ОТЧЁТ", bold=True)
            add_paragraph(doc, "по лабораторной работе")
    else:
        add_paragraph(doc, "Пояснительная записка", bold=True)
        add_paragraph(doc, "к курсовому проекту")

    add_paragraph(doc, "на тему:")
    add_empty_lines(doc, 2)

    topic_line = topic if topic.startswith("«") else f"«{topic}»"
    add_paragraph(doc, topic_line, bold=True)
    add_empty_lines(doc, 1)
    add_paragraph(doc, doc_code)
    if variant is not None:
        add_paragraph(doc, f"Вариант {variant}")
    add_empty_lines(doc, 1)

    if kind == "lab":
        _signature_table(
            doc,
            [
                (f"Выполнил студент группы {student_group}", WD_ALIGN_PARAGRAPH.LEFT, False),
                (student_name, WD_ALIGN_PARAGRAPH.LEFT, False),
                ("______________________________", None, False),
                ("(подпись студента)", WD_ALIGN_PARAGRAPH.CENTER, False),
            ],
        )
    else:
        _signature_table(
            doc,
            [
                (f"Выполнил студент группы {student_group}", WD_ALIGN_PARAGRAPH.LEFT, False),
                (student_name, WD_ALIGN_PARAGRAPH.LEFT, False),
                ("______________________________", None, False),
                ("(подпись студента)", WD_ALIGN_PARAGRAPH.CENTER, False),
                (
                    "Курсовой проект представлен на проверку ___.___.2026",
                    WD_ALIGN_PARAGRAPH.LEFT,
                    False,
                ),
                ("______________________________", None, False),
                ("(подпись студента)", WD_ALIGN_PARAGRAPH.CENTER, False),
            ],
        )

    add_empty_lines(doc, 5)
    add_paragraph(doc, city_year)

def add_toc_entry(doc, title: str, page: int | str, *, level: int = 0):
    """Строка содержания: отточие через tab leader, номер страницы справа (СТП 2.2.7)."""
    indent = "\u00a0\u00a0" * level
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    run = p.add_run(f"{indent}{title}\t{page}")
    set_run_font(run)


def build_toc(doc, entries: list[tuple[str, int, int | str]]):
    """
    entries: (заголовок, уровень_вложенности, номер_страницы)
    level 0 — раздел, 1 — подраздел (смещение на 2 знака по СТП).
    """
    p = doc.add_paragraph()
    run = p.add_run("СОДЕРЖАНИЕ")
    set_run_font(run, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(18)
    add_empty_lines(doc, 1)
    for title, level, page in entries:
        add_toc_entry(doc, title, page, level=level)
