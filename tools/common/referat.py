"""Реферат курсового проекта по образцу ПСП (СТП БГУИР)."""

from __future__ import annotations

from dataclasses import dataclass

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

from tools.common.docx_stp import FONT, INDENT, SIZE

FONT_NAME = FONT


@dataclass
class ReferatConfig:
    doc_code: str = "БГУИР КП 6-05-0611-01 010 ПЗ"
    author_bib: str = "Демидов, И. С."
    author_short: str = "И. С. Демидов"
    topic_title: str = "Название темы курсового проекта"
    keywords_caps: str = "КЛЮЧЕВОЕ, СЛОВО, ТЕХНОЛОГИЯ"
    originality_pct: int = 95
    pages: int = 50
    figures: int = 20
    tables: int = 5
    sources: int = 10
    appendices: int = 3
    goal: str = "цель проектирования — описать в content.py"
    methodology: str = "методология — описать в content.py"
    results: str = "результаты — описать в content.py"
    tech_stack: str = "стек технологий — описать в content.py"
    application: str = "область применения — описать в content.py"
    year: int = 2026


def _run(p, text: str, *, bold: bool = False, italic: bool = False):
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = SIZE
    r.bold = bold
    r.italic = italic
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    return r


def _fmt(p, *, indent: bool = True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    pf = p.paragraph_format
    pf.first_line_indent = INDENT if indent else Cm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(18)
    p.alignment = align


def stats_line(cfg: ReferatConfig) -> str:
    return (
        f"Пояснительная записка {cfg.pages} с., {cfg.figures} рис., {cfg.tables} табл., "
        f"{cfg.sources} источников, {cfg.appendices} приложения."
    )


def antiplagiat_intro_text(cfg: ReferatConfig) -> str:
    return (
        f"Курсовая работа выполнена самостоятельно, проверена в системе «Антиплагиат». "
        f"Процент оригинальности составляет {cfg.originality_pct}%. Цитирования обозначены "
        f"ссылками на публикации, указанными в «Списке использованных источников». "
        f"Скриншот приведён в приложении А (рисунок А.1)."
    )


def _blank_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def append_referat_body(anchor: Paragraph, cfg: ReferatConfig) -> Paragraph:
    last = anchor
    for block in (
        lambda p: (_run(p, cfg.doc_code), _fmt(p, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)),
        lambda p: (
            _run(p, cfg.author_bib, bold=True),
            _run(p, f" {cfg.topic_title} / {cfg.author_short}. – Минск: БГУИР, {cfg.year}. –"),
            _fmt(p, indent=True),
        ),
        lambda p: (_run(p, stats_line(cfg)), _fmt(p, indent=True, align=WD_ALIGN_PARAGRAPH.LEFT)),
        lambda p: (_run(p, cfg.keywords_caps), _fmt(p, indent=True)),
    ):
        last = _blank_after(last)
        p = _blank_after(last)
        block(p)
        last = p

    for label, text in (
        (["Цель ", "проектирования"], f": {cfg.goal}"),
        (["Методология проведения работы"], f": {cfg.methodology}"),
        (["Результаты работы"], f": {cfg.results}"),
        (["Область применения результатов"], f": {cfg.application}"),
    ):
        last = _blank_after(last)
        p = _blank_after(last)
        for part in label:
            _run(p, part, italic=True)
        _run(p, text)
        _fmt(p, indent=True)
        last = p

    last = _blank_after(last)
    p = _blank_after(last)
    _run(p, cfg.tech_stack)
    _fmt(p, indent=True)
    return p


def insert_antiplagiat_before_section1(doc: Document, cfg: ReferatConfig) -> None:
    text = antiplagiat_intro_text(cfg)
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("1 ") and "АНАЛИЗ" in paragraph.text.upper():
            new_p = OxmlElement("w:p")
            paragraph._element.addprevious(new_p)
            para = Paragraph(new_p, paragraph._parent)
            _run(para, text)
            _fmt(para, indent=True)
            _blank_after(para)
            return
    raise ValueError("Не найден раздел 1 в документе")
